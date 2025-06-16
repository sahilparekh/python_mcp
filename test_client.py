#!/usr/bin/env python3
"""
Test client for the FastMCP Python Code Executor server.
This demonstrates how to interact with the MCP server programmatically.
"""

import asyncio
from fastmcp import Client

async def test_python_executor():
    """Test the Python code executor MCP server."""
    
    # Connect to the server using the in-memory transport
    # This works by importing the server directly
    from utils import mcp
    
    async with Client(mcp) as client:
        print("🚀 Connected to FastMCP Python Code Executor")
        print("=" * 50)
        
        # Test 1: Get allowed libraries documentation via resource
        print("\n📚 Test 1: Getting allowed libraries documentation via resource...")
        try:
            libs_resource = await client.read_resource("docs://allowed-libraries")
            print(f"Allowed libraries resource: {libs_resource.text}")
        except Exception as e:
            print(f"Resource read failed: {e}")
        
        # Test 2: Get specific library documentation
        print("\n🔍 Test 2: Getting specific library documentation...")
        try:
            pandas_docs = await client.read_resource("docs://library/pandas")
            print(f"Pandas documentation: {pandas_docs.text}")
        except Exception as e:
            print(f"Library-specific resource read failed: {e}")
        
        # Test 3: Get server documentation
        print("\n📋 Test 3: Getting server documentation...")
        try:
            server_docs = await client.read_resource("docs://server-info")
            print(f"Server documentation: {server_docs.text}")
        except Exception as e:
            print(f"Server documentation read failed: {e}")
        
        # Test 4: Get server information via tool
        print("\n📋 Test 4: Getting server information via tool...")
        info_result = await client.call_tool("get_server_info", {})
        print(f"Server info: {info_result.text}")
        
        # Test 5: Validate code with allowed libraries
        print("\n✅ Test 5: Validating code with allowed libraries...")
        valid_code = """
import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt

# This should pass validation
print("Using allowed libraries!")
"""
        validation_result = await client.call_tool("validate_code_libraries", {"code": valid_code})
        print(f"Validation result: {validation_result.text}")
        
        # Test 6: Validate code with disallowed libraries
        print("\n❌ Test 6: Validating code with disallowed libraries...")
        invalid_code = """
import numpy as np  # Not in allowed list
import requests     # Not in allowed list
print("This should fail validation")
"""
        invalid_validation = await client.call_tool("validate_code_libraries", {"code": invalid_code})
        print(f"Invalid validation result: {invalid_validation.text}")
        
        # Test 7: Execute valid Python code
        print("\n🐍 Test 7: Executing valid Python code...")
        simple_code = """
import pandas as pd
import matplotlib.pyplot as plt
import json

print("Hello from FastMCP with allowed libraries!")

# Create some sample data using pandas
data = {'Name': ['Alice', 'Bob', 'Charlie'], 'Age': [25, 30, 35]}
df = pd.DataFrame(data)
print("Created DataFrame:")
print(df)

# Save as JSON
df.to_json('sample_data.json', orient='records', indent=2)

# Create a simple plot
plt.figure(figsize=(8, 6))
plt.bar(df['Name'], df['Age'])
plt.title('Sample Data Visualization')
plt.ylabel('Age')
plt.savefig('sample_plot.png')
plt.close()

print("Generated sample_data.json and sample_plot.png")
"""
        exec_result = await client.call_tool("execute_python_code", {"code": simple_code})
        print(f"Execution result: {exec_result.text}")
        
        # Test 8: Try to execute code with disallowed libraries
        print("\n🚫 Test 8: Attempting to execute code with disallowed libraries...")
        disallowed_code = """
import numpy as np
import requests
print("This should be blocked!")
"""
        blocked_result = await client.call_tool("execute_python_code", {"code": disallowed_code})
        print(f"Blocked execution result: {blocked_result.text}")
        
        # Test 9: Execute code that generates multiple files with allowed libraries
        print("\n📁 Test 9: Executing code that generates multiple files...")
        multi_file_code = """
import pandas as pd
import json
from docx import Document
from reportlab.pdfgen import canvas

# Create a JSON file
data = {"message": "Hello", "timestamp": "2025-06-15", "version": "2.0"}
with open('data.json', 'w') as f:
    json.dump(data, f, indent=2)

# Create a CSV file using pandas
df = pd.DataFrame({
    'Name': ['Alice', 'Bob', 'Charlie'],
    'Age': [30, 25, 35],
    'City': ['New York', 'San Francisco', 'Chicago']
})
df.to_csv('data.csv', index=False)

# Create a Word document
doc = Document()
doc.add_heading('Sample Report', 0)
doc.add_paragraph('This is a sample document created with python-docx.')
doc.save('report.docx')

# Create a simple PDF
c = canvas.Canvas('report.pdf')
c.drawString(100, 750, "Sample PDF created with ReportLab")
c.save()

print("Generated JSON, CSV, DOCX, and PDF files using allowed libraries")
print(f"Files created: {sorted([f for f in os.listdir('.') if f != 'script.py'])}")
"""
        multi_result = await client.call_tool("execute_python_code", {"code": multi_file_code})
        print(f"Multi-file execution result: {multi_result.text}")
        
        # Test 10: List artifacts (this might fail if GCS is not configured)
        print("\n📦 Test 10: Listing recent artifacts...")
        try:
            artifacts_result = await client.call_tool("list_bucket_artifacts", {"limit": 5})
            print(f"Artifacts: {artifacts_result.text}")
        except Exception as e:
            print(f"Note: Artifact listing failed (GCS not configured): {e}")
        
        # Test 11: Execute code with error handling
        print("\n❌ Test 11: Testing error handling...")
        error_code = """
import pandas as pd
print("This will cause an error...")
raise ValueError("Intentional error for testing")
"""
        error_result = await client.call_tool("execute_python_code", {"code": error_code})
        print(f"Error handling result: {error_result.text}")
        
    print("\n✅ All tests completed!")

if __name__ == "__main__":
    print("FastMCP Python Code Executor - Test Client")
    print("This script tests the MCP server functionality")
    print()
    
    # Run the async test
    asyncio.run(test_python_executor())
